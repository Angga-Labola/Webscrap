"""Document and URL content extractors for ingestion."""

import os
import re
import urllib.request
import urllib.parse
from pathlib import Path
from typing import Optional, Dict, Any
from bs4 import BeautifulSoup

try:
    import pypdf
except ImportError:
    pypdf = None

try:
    import docx
except ImportError:
    docx = None

try:
    from pptx import Presentation as PptxPresentation
except ImportError:
    PptxPresentation = None


def extract_document_text(file_path: str, original_filename: Optional[str] = None) -> Dict[str, Any]:
    """Extracts raw text content from uploaded files (PDF, DOCX, TXT, MD, PPTX)."""
    p = Path(file_path)
    if not p.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    filename = (original_filename or p.name).lower()
    text = ""
    doc_type = "unknown"

    try:
        if filename.endswith(".pdf"):
            doc_type = "pdf"
            if pypdf:
                reader = pypdf.PdfReader(str(p))
                pages = []
                for i, page in enumerate(reader.pages):
                    page_text = page.extract_text() or ""
                    if page_text.strip():
                        pages.append(f"--- Page {i+1} ---\n{page_text.strip()}")
                text = "\n\n".join(pages)
            else:
                text = "[pypdf library not available to extract PDF content]"

        elif filename.endswith(".docx"):
            doc_type = "docx"
            if docx:
                doc = docx.Document(str(p))
                paragraphs = [para.text.strip() for para in doc.paragraphs if para.text.strip()]
                # Also include table text
                for table in doc.tables:
                    for row in table.rows:
                        row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                        if row_text:
                            paragraphs.append(row_text)
                text = "\n\n".join(paragraphs)
            else:
                text = "[python-docx library not available to extract DOCX content]"

        elif filename.endswith(".pptx"):
            doc_type = "pptx"
            if PptxPresentation:
                prs = PptxPresentation(str(p))
                slides_text = []
                for idx, slide in enumerate(prs.slides, start=1):
                    shapes_text = []
                    for shape in slide.shapes:
                        if shape.has_text_frame:
                            for paragraph in shape.text_frame.paragraphs:
                                if paragraph.text.strip():
                                    shapes_text.append(paragraph.text.strip())
                    if shapes_text:
                        slides_text.append(f"--- Slide {idx} ---\n" + "\n".join(shapes_text))
                text = "\n\n".join(slides_text)
            else:
                text = "[python-pptx not available]"

        elif any(filename.endswith(ext) for ext in [".txt", ".md", ".markdown", ".csv", ".json", ".log"]):
            doc_type = "text"
            with open(p, "r", encoding="utf-8", errors="replace") as f:
                text = f.read()

        else:
            # Fallback to text reading
            doc_type = "text_fallback"
            with open(p, "r", encoding="utf-8", errors="replace") as f:
                text = f.read()

    except Exception as e:
        text = f"[Error extracting text from {filename}: {str(e)}]"

    # Clean text
    clean_text = re.sub(r"\n{3,}", "\n\n", text).strip()
    return {
        "filename": original_filename or p.name,
        "type": doc_type,
        "content": clean_text,
        "length": len(clean_text)
    }


def extract_url_text(url: str, max_chars: int = 15000) -> Dict[str, Any]:
    """Fetches web page content, parses HTML, and extracts clean article text."""
    clean_url = url.strip()
    if not clean_url.startswith(("http://", "https://")):
        clean_url = "https://" + clean_url

    req = urllib.request.Request(
        clean_url,
        headers={
            "User-Agent": (
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) "
                "Chrome/124.0.0.0 Safari/537.36"
            ),
            "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
            "Accept-Language": "en-US,en;q=0.9",
        }
    )

    try:
        with urllib.request.urlopen(req, timeout=12) as response:
            html = response.read().decode("utf-8", errors="replace")

        soup = BeautifulSoup(html, "html.parser")

        # Extract title
        title = ""
        if soup.title and soup.title.string:
            title = soup.title.string.strip()

        # Remove irrelevant and clutter tags
        for tag in soup(["script", "style", "nav", "footer", "header", "aside", "noscript", "svg", "form"]):
            tag.decompose()

        # Prioritize main content elements if present
        main_content = soup.find("article") or soup.find("main") or soup.find("div", class_=re.compile(r"content|article|post|body", re.I)) or soup.body

        if main_content:
            text = main_content.get_text(separator="\n", strip=True)
        else:
            text = soup.get_text(separator="\n", strip=True)

        # Collapse whitespace
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        cleaned = "\n".join(lines)

        # Truncate if excessively long to preserve token budget
        if len(cleaned) > max_chars:
            cleaned = cleaned[:max_chars] + f"\n\n[Content truncated after {max_chars} characters...]"

        return {
            "url": clean_url,
            "title": title,
            "content": cleaned,
            "length": len(cleaned),
            "success": True,
            "error": None
        }

    except Exception as e:
        return {
            "url": clean_url,
            "title": "",
            "content": "",
            "length": 0,
            "success": False,
            "error": str(e)
        }

